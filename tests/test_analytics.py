"""Spike Analytics (C7): small datasets whose answers were computed by hand, the sandbox that keeps SQL inside the
loaded tables, charts built by SPIKE, data_python on copies, the chat showing query, table and chart, and the
report export."""

import base64
import io
import json
import os
import sqlite3

import pytest
from fastapi.testclient import TestClient

from spike.analytics import engine
from spike.analytics.charts import build_chart
from spike.analytics.engine import DataError, Result, Session, describe, run_query
from spike.analytics.tools import DataChart, DataDescribe, DataPython, DataSql
from spike.server.app import create_app
from spike.server.state import AppState
from spike.tools.builtin import Workspace
from tests.helpers import ScriptedProvider, call, text

duckdb = pytest.importorskip("duckdb", reason="the data extra (duckdb, pandas)")
pytest.importorskip("pandas")

# revenue = unidades × precio: 25, 40, 60, 20, 5, 10 → total 160; norte 90, sur 70; Jan 65, Feb 85, Mar 10
SALES = """fecha,region,producto,unidades,precio
2024-01-15,norte,A,10,2.50
2024-01-20,sur,B,4,10.00
2024-02-03,norte,B,6,10.00
2024-02-10,sur,A,8,2.50
2024-02-28,norte,A,2,2.50
2024-03-05,sur,B,1,10.00
"""


@pytest.fixture
def ws(tmp_path) -> Workspace:
    root = tmp_path / "ws"
    root.mkdir()
    (root / "ventas 2024.csv").write_text(SALES, encoding="utf-8")
    (root / "importes.csv").write_text("mes;importe\nene;10,5\nfeb;20,25\n", encoding="utf-8")  # Spanish Excel
    (root / "eventos.jsonl").write_text('{"cat": "x", "v": 1}\n{"cat": "x", "v": 3}\n{"cat": "y", "v": 10}\n')
    duckdb.sql(f"COPY (SELECT range + 1 AS n FROM range(10)) TO '{root / 'nums.parquet'}' (FORMAT parquet)")
    db = sqlite3.connect(root / "tienda.sqlite")
    db.executescript("CREATE TABLE clientes (id INTEGER, nombre TEXT); CREATE TABLE pedidos (id INTEGER, "
                     "cliente_id INTEGER, total REAL); INSERT INTO clientes VALUES (1, 'Ana'), (2, 'Luis'); "
                     "INSERT INTO pedidos VALUES (1, 1, 30.0), (2, 1, 12.5), (3, 2, 7.0);")
    db.commit()
    db.close()
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    first = book.active
    first.title = "2023"
    for row in (("región", "ventas"), ("norte", 100), ("sur", 50)):
        first.append(row)
    second = book.create_sheet("2024")
    for row in (("región", "ventas"), ("norte", 120), ("sur", 80)):
        second.append(row)
    book.save(root / "libro.xlsx")
    return Workspace(root)


def q(ws: Workspace, files: list[str], sql: str, limit: int = 200) -> list[list]:
    return run_query(ws, files, sql, limit).rows


def test_csv_answers_match_the_hand_computed_values(ws):
    files = ["ventas 2024.csv"]
    assert q(ws, files, "SELECT sum(unidades * precio) FROM ventas_2024") == [[160.0]]
    assert q(ws, files, "SELECT region, sum(unidades * precio) AS ingresos FROM ventas_2024 GROUP BY region "
                        "ORDER BY region") == [["norte", 90.0], ["sur", 70.0]]
    assert q(ws, files, "SELECT strftime(fecha, '%Y-%m') AS mes, sum(unidades * precio) FROM ventas_2024 "
                        "GROUP BY 1 ORDER BY 1") == [["2024-01", 65.0], ["2024-02", 85.0], ["2024-03", 10.0]]
    assert q(ws, files, "SELECT avg(precio), sum(unidades), count(*) FROM ventas_2024") == [[6.25, 31, 6]]
    result = run_query(ws, files, "SELECT * FROM ventas_2024", 200)
    assert result.types[0] == "DATE" and result.rows[0][0] == "2024-01-15"  # dates come back as ISO text
    assert result.tables[0].name == "ventas_2024" and result.tables[0].rows == 6


def test_decimal_commas_semicolons_json_parquet_excel_and_sqlite(ws):
    assert q(ws, ["importes.csv"], "SELECT sum(importe) FROM importes") == [[30.75]]
    assert q(ws, ["eventos.jsonl"], "SELECT cat, avg(v) FROM eventos GROUP BY cat ORDER BY cat") == [["x", 2.0],
                                                                                                    ["y", 10.0]]
    assert q(ws, ["nums.parquet"], "SELECT sum(n), avg(n) FROM nums") == [[55, 5.5]]
    excel = run_query(ws, ["libro.xlsx"], 'SELECT a."región", b.ventas - a.ventas AS crece FROM libro_2023 a '
                                          'JOIN libro_2024 b USING ("región") ORDER BY 1', 200)
    assert excel.rows == [["norte", 20], ["sur", 30]]
    assert [(t.name, t.sheet) for t in excel.tables] == [("libro_2023", "2023"), ("libro_2024", "2024")]
    assert q(ws, ["tienda.sqlite"], "SELECT c.nombre, sum(p.total) FROM pedidos p JOIN clientes c ON "
                                    "c.id = p.cliente_id GROUP BY 1 ORDER BY 1") == [["Ana", 42.5], ["Luis", 7.0]]
    both = run_query(ws, ["ventas 2024.csv", "nums.parquet"], "SELECT (SELECT count(*) FROM ventas_2024) + "
                                                              "(SELECT max(n) FROM nums)", 10)
    assert both.rows == [[16]] and [t.name for t in both.tables] == ["ventas_2024", "nums"]


def test_describe_gives_columns_statistics_and_first_rows(ws):
    [(table, summary, sample)] = describe(ws, ["ventas 2024.csv"])
    assert table.columns[3] == ("unidades", "BIGINT") and table.rows == 6
    stats = {row[0]: dict(zip(summary.columns, row, strict=True)) for row in summary.rows}
    assert stats["unidades"]["min"] == "1" and stats["unidades"]["max"] == "10"
    assert float(stats["precio"]["avg"]) == 6.25 and stats["region"]["approx_unique"] == 2
    assert len(sample.rows) == 5


@pytest.mark.parametrize("sql", [
    "SELECT * FROM read_csv('/etc/passwd')", "SELECT * FROM '/etc/hostname'", "COPY ventas_2024 TO 'out.csv'",
    "ATTACH 'other.db'", "INSTALL httpfs", "LOAD sqlite", "SET enable_external_access = true",
    "RESET lock_configuration", "SELECT getenv('HOME')", "SET memory_limit = '50GB'",
    "SELECT * FROM read_text('ventas 2024.csv')"])
def test_sql_cannot_leave_the_loaded_tables(ws, sql):
    with pytest.raises(DataError, match="SQL error"):
        run_query(ws, ["ventas 2024.csv"], sql, 10)
    assert not (ws.root / "out.csv").exists()


def test_only_workspace_files_load_and_bad_input_is_explained(ws, tmp_path):
    (tmp_path / "secret.csv").write_text("a\n1\n")
    with pytest.raises(PermissionError):
        run_query(ws, ["../secret.csv"], "SELECT 1", 1)
    os.symlink(tmp_path / "secret.csv", ws.root / "link.csv")
    with pytest.raises(PermissionError):  # a link inside the workspace that points outside it
        run_query(ws, ["link.csv"], "SELECT 1", 1)
    with pytest.raises(DataError, match="not a file"):
        run_query(ws, ["missing.csv"], "SELECT 1", 1)
    (ws.root / "notes.docx").write_bytes(b"x")
    with pytest.raises(DataError, match="unsupported"):
        run_query(ws, ["notes.docx"], "SELECT 1", 1)
    (ws.root / "fake.sqlite").write_text("not a database")
    with pytest.raises(DataError, match="not a SQLite"):
        run_query(ws, ["fake.sqlite"], "SELECT 1", 1)
    with pytest.raises(DataError, match="SQL error: .*ventas"):  # DuckDB's own hint, so the model can fix it
        run_query(ws, ["ventas 2024.csv"], "SELECT nope FROM ventas_2024", 1)


def test_a_long_query_is_stopped_and_results_are_capped(ws, monkeypatch):
    monkeypatch.setattr(engine, "TIMEOUT_S", 0.5)
    with pytest.raises(DataError, match="longer than"):
        run_query(ws, ["nums.parquet"], "SELECT count(*) FROM range(100000000000) a", 1)
    capped = run_query(ws, ["nums.parquet"], "SELECT n FROM nums ORDER BY n", 3)
    assert capped.rows == [[1], [2], [3]] and capped.truncated
    assert "(first 3 rows; there are more)" in capped.text()


def test_a_session_is_locked_after_loading(ws):
    session = Session(ws, ["ventas 2024.csv"])
    try:
        session.load()
        assert session.con.execute("SELECT current_setting('enable_external_access')").fetchone() == (False,)
        assert session.con.execute("SELECT current_setting('lock_configuration')").fetchone() == (True,)
    finally:
        session.close()


def result(columns, types, rows) -> Result:
    return Result(columns, types, rows, False)


def test_charts_are_built_by_spike_with_inline_data():
    r = result(["region", "ingresos"], ["VARCHAR", "DOUBLE"], [["norte", 90.0], ["sur", 70.0]])
    spec = build_chart(r, "bar", title="Ingresos")
    assert spec["data"] == {"values": [{"x": "norte", "y": 90.0}, {"x": "sur", "y": 70.0}]}
    assert spec["encoding"]["x"] == {"field": "x", "type": "nominal", "title": "region", "sort": None,
                                     "axis": {"labelAngle": 0, "labelLimit": 160}}
    long = result(["producto", "n"], ["VARCHAR", "BIGINT"], [[f"Producto con nombre largo {i}", i] for i in range(3)])
    assert build_chart(long)["encoding"]["x"]["axis"]["labelAngle"] == -40  # long labels slant instead of overlapping
    assert spec["encoding"]["y"]["title"] == "ingresos" and spec["title"] == "Ingresos"
    assert "url" not in json.dumps(spec["data"]) and "href" not in json.dumps(spec)
    dates = result(["dia", "n"], ["DATE", "BIGINT"], [["2024-01-01", 3], ["2024-01-02", 5]])
    line = build_chart(dates, "line")
    assert line["encoding"]["x"]["type"] == "temporal" and line["encoding"]["x"]["scale"] == {"type": "utc"}
    years = build_chart(result(["año", "n"], ["INTEGER", "BIGINT"], [[2023, 1], [2024, 2]]), "bar")
    assert years["encoding"]["x"]["type"] == "ordinal"  # one bar per year
    wide = result(["mes", "ventas", "costes"], ["VARCHAR", "DOUBLE", "DOUBLE"], [["ene", 10, 4], ["feb", 12, 5]])
    melted = build_chart(wide, "line", y=["ventas", "costes"])
    assert melted["data"]["values"][:2] == [{"x": "ene", "y": 10, "series": "ventas"},
                                            {"x": "ene", "y": 4, "series": "costes"}]
    pie = build_chart(r, "arc")
    assert set(pie["encoding"]) == {"theta", "color", "tooltip"}


@pytest.mark.parametrize("kwargs, message", [
    ({"mark": "image"}, "mark must be"), ({"y": "region"}, "y must be numeric"), ({"x": "nope"}, "x must be one"),
    ({"mark": "arc", "color": "region"}, "pie chart")])
def test_bad_charts_are_refused(kwargs, message):
    r = result(["region", "ingresos"], ["VARCHAR", "DOUBLE"], [["norte", 90.0]])
    with pytest.raises(DataError, match=message):
        build_chart(r, **kwargs)
    with pytest.raises(DataError, match="aggregate first"):
        build_chart(Result(["a", "b"], ["VARCHAR", "DOUBLE"], [["x", 1]], True))


async def test_the_tools_return_the_query_the_table_and_the_chart(ws):
    sql = DataSql(ws)
    out = await sql.run({"sql": "SELECT region, sum(unidades * precio) AS ingresos FROM ventas_2024 GROUP BY 1 "
                                "ORDER BY 1", "files": ["ventas 2024.csv"], "title": "Ingresos por región"})
    assert out.ok and "norte | 90" in out.content and "ventas_2024 from ventas 2024.csv: 6 rows" in out.content
    assert out.private == [out.content]  # the user's data: withheld from external providers unless allowed
    ui = out.data["ui"]
    assert ui["kind"] == "data_table" and ui["rows"] == [["norte", 90.0], ["sur", 70.0]] and "GROUP BY" in ui["sql"]
    chart = await DataChart(ws).run({"sql": "SELECT region, sum(unidades) u FROM ventas_2024 GROUP BY 1 ORDER BY 1",
                                     "files": "ventas 2024.csv", "mark": "bar"})
    assert chart.ok and chart.data["ui"]["spec"]["data"]["values"] == [{"x": "norte", "y": 18}, {"x": "sur", "y": 13}]
    bad = await DataChart(ws).run({"sql": "SELECT region FROM ventas_2024", "files": ["ventas 2024.csv"]})
    assert not bad.ok and "at least two columns" in bad.content
    desc = await DataDescribe(ws).run({"files": ["libro.xlsx"]})
    assert desc.ok and [i["title"] for i in desc.data["ui"]["items"]] == ["libro_2023", "libro_2024"]
    assert '"región" VARCHAR' in desc.content  # how the model must write that column
    missing = await DataSql(ws).run({"sql": "SELECT 1"})
    assert not missing.ok and "give the data files" in missing.content


async def test_data_python_runs_on_copies(ws):
    tool = DataPython(ws)
    assert tool.default_policy == "ask"
    code = ("import pandas as pd\ndf = pd.read_csv('data/ventas 2024.csv')\n"
            "print((df.unidades * df.precio).groupby(df.region).sum().to_dict())\n"
            "open('data/ventas 2024.csv', 'w').write('borrado')")
    out = await tool.run({"code": code, "files": ["ventas 2024.csv"]})
    assert out.ok, out.content
    assert out.content == "{'norte': 90.0, 'sur': 70.0}"
    assert (ws.root / "ventas 2024.csv").read_text(encoding="utf-8") == SALES  # the original is untouched
    assert out.data["ui"] == {"kind": "code_run", "language": "python", "code": code, "output": out.content,
                              "files": ["ventas 2024.csv"], "returncode": 0}
    outside = await tool.run({"code": "print(1)", "files": ["../x.csv"]})
    assert not outside.ok and "outside the workspace" in outside.content
    plain = await tool.run({"code": "print(2 + 2)"})
    assert plain.content == "4" and plain.private == []


# ---------------------------------------------------------------- the app
@pytest.fixture
def env(tmp_path, monkeypatch):
    for name, value in (("SPIKE_HOME", tmp_path / "home"), ("SPIKE_WORKSPACE", tmp_path / "ws"),
                        ("SPIKE_NO_KEYRING", "1")):
        monkeypatch.setenv(name, str(value))
    local = ScriptedProvider([], name="spike-local")

    def factory(cfg):
        local.config = cfg
        return local

    state = AppState(home=tmp_path / "home", provider_factory=factory)
    return TestClient(create_app(state)), local


def events_of(response) -> list[dict]:
    return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]


def test_a_project_chat_answers_from_its_files_with_a_chart(env):
    client, local = env
    project = client.post("/api/projects", json={"name": "Ventas"}).json()
    client.put(f"/api/projects/{project['id']}/files/ventas.csv", content=SALES.encode())
    sql = "SELECT region, sum(unidades * precio) AS ingresos FROM ventas GROUP BY 1 ORDER BY 2 DESC"
    local._script = [call("data_chart", {"sql": sql, "files": ["ventas.csv"], "mark": "bar", "title": "Ingresos"}),
                     text("El norte ingresa 90 y el sur 70.")]
    r = client.post("/api/chat", json={"message": "¿Qué región vende más? Haz un gráfico", "mode": "data",
                                       "project_id": project["id"]})
    result = next(e for e in events_of(r) if e["type"] == "tool_result")
    assert result["ok"] and result["ui"]["kind"] == "chart" and result["ui"]["sql"] == sql
    assert result["ui"]["spec"]["data"]["values"] == [{"x": "norte", "y": 90.0}, {"x": "sur", "y": 70.0}]
    names = {t.name for t in local.requests[0].tools}
    assert {"data_describe", "data_sql", "data_chart", "data_python"} <= names and "python" not in names
    assert "never from your own arithmetic" in local.requests[0].messages[0].content
    conv = events_of(r)[0]["id"]
    stored = client.get(f"/api/conversations/{conv}").json()
    assert any(t.get("ui", {}).get("kind") == "chart" for m in stored["messages"] for t in m["meta"].get("tools", []))


def png(width=40, height=20) -> str:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), "#5FBB47").save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


BLOCKS = [{"kind": "text", "title": "Pregunta", "text": "¿Qué región vende más?"},
          {"kind": "text", "text": "El **norte**: 90 frente a 70."},
          {"kind": "code", "title": "Consulta", "language": "sql", "text": "SELECT región, sum(x) FROM ventas"},
          {"kind": "table", "columns": ["región", "ingresos"], "rows": [["norte", 90], ["sur", None]]},
          {"kind": "image", "title": "Ingresos", "png": png()}]


def test_the_report_exports_to_word_and_pdf(env):
    pytest.importorskip("docx")
    pytest.importorskip("reportlab")
    client, _ = env
    r = client.post("/api/analytics/report", json={"format": "docx", "title": "Ventas por región", "blocks": BLOCKS})
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"] == 'attachment; filename="ventas-por-region.docx"'
    import docx

    doc = docx.Document(io.BytesIO(r.content))
    texts = [p.text for p in doc.paragraphs]
    assert "Ventas por región" in texts and "SELECT región, sum(x) FROM ventas" in texts
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["norte", "90"] and doc.tables[0].rows[2].cells[1].text == ""
    assert len(doc.inline_shapes) == 1
    pdf = client.post("/api/analytics/report", json={"format": "pdf", "title": "Ventas", "blocks": BLOCKS})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and pdf.headers["content-type"] == "application/pdf"
    from pypdf import PdfReader

    page = PdfReader(io.BytesIO(pdf.content)).pages[0].extract_text()
    assert "norte" in page and "SELECT región" in page


@pytest.mark.parametrize("body, message", [
    ({"format": "html", "blocks": BLOCKS}, "format"), ({"format": "pdf", "blocks": []}, "1 to"),
    ({"format": "pdf", "blocks": [{"kind": "image", "png": base64.b64encode(b"GIF89a").decode()}]}, "PNG"),
    ({"format": "pdf", "blocks": [{"kind": "script", "text": "x"}]}, "unknown block"),
    ({"format": "pdf", "blocks": [{"kind": "image", "png": png(5000, 10)}]}, "pixels")])
def test_bad_reports_are_refused(env, body, message):
    client, _ = env
    r = client.post("/api/analytics/report", json=body)
    assert r.status_code == 400 and message in r.json()["detail"]
