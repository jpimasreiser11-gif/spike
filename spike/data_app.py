"""Bounded local table apps. Only cell edits are supported; headers, sheet structure and formulas are read-only.

XLSX uses the already declared optional openpyxl extra. Workbooks are rewritten by openpyxl, so advanced Excel
features unsupported by that library may not survive; this template is intended for small plain data tables.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import tempfile
import zipfile
from decimal import Decimal
from pathlib import Path

from spike.projects import ProjectFiles, clean_path

MAX_BYTES = 2 * 1024**2


def safe_path(root: Path, relative: str) -> Path:
    path = root.joinpath(*clean_path(relative).parts)
    if root.resolve() != root or ProjectFiles._linked(root, path) or not path.resolve().is_relative_to(root.resolve()):
        raise ValueError("data_app_path")
    return path


def digest(path: Path) -> str:
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("data_app_limit")
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook(path: Path):
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("data_app_requires_xlsx") from exc
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 20 * MAX_BYTES:
            raise ValueError("data_app_limit")
    return openpyxl.load_workbook(path, keep_links=False)


def table(path: Path) -> list[list[str]]:
    digest(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = []
            cells = 0
            for row in csv.reader(stream):
                cells += len(row)
                if len(rows) >= 2001 or len(row) > 100 or cells > 20000:
                    raise ValueError("data_app_limit")
                rows.append(row)
    elif path.suffix.lower() == ".xlsx":
        book = workbook(path)
        try:
            sheet = book.worksheets[0]
            if sheet.max_row > 2001 or sheet.max_column > 100:
                raise ValueError("data_app_limit")
            rows = [["" if c.value is None else str(c.value) for c in row] for row in sheet]
        finally:
            book.close()
    else:
        raise ValueError("data_app_format")
    if not rows or not rows[0] or len(rows) > 2001 or len(rows[0]) > 100 or sum(map(len, rows)) > 20000:
        raise ValueError("data_app_limit")
    if any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("data_app_table")
    return rows


def document(rows: list[list[str]], lang: str, title: str = "", labels: list[str] | None = None) -> str:
    labels = labels or (["Buscar", "Proponer cambios", "Los cambios requieren aprobación en SPIKE."] if lang == "es" else
              ["Search", "Propose changes", "Changes require approval in SPIKE."])
    payload = json.dumps({"rows": rows, "labels": labels, "title": title}, ensure_ascii=True).replace("<", "\\u003c")
    return '''<style>body{font:16px system-ui;padding:20px}input{padding:6px;max-width:180px}td,th{padding:4px}table{border-collapse:collapse}button{padding:10px}</style>
<h1 id="title"></h1><label id="searchLabel"><input id="search" type="search"></label><p id="hint"></p><table><thead></thead><tbody></tbody></table><button id="propose"></button>
<script>
const data = PAYLOAD;
const rows = data.rows, edits = new Map();
document.querySelector('#title').textContent = data.title;
const search = document.querySelector('#search');
document.querySelector('#searchLabel').prepend(data.labels[0] + ' ');
document.querySelector('#hint').textContent = data.labels[2];
document.querySelector('#propose').textContent = data.labels[1];
const head = document.createElement('tr');
rows[0].forEach(text => { const th = document.createElement('th'); th.textContent = text; head.append(th); });
document.querySelector('thead').append(head);
function render() {
 const body = document.querySelector('tbody'); body.replaceChildren();
 rows.slice(1).forEach((row, r) => {
  if (!row.join(' ').toLowerCase().includes(search.value.toLowerCase())) return;
  const tr = document.createElement('tr');
  row.forEach((text, c) => {
   const td = document.createElement('td'), input = document.createElement('input');
   input.value = edits.get(r + ':' + c)?.value ?? text;
   input.setAttribute('aria-label', rows[0][c] + ' ' + (r + 1));
   input.oninput = () => { const key = r + ':' + c;
    if (input.value === text) edits.delete(key); else edits.set(key, {row:r, column:c, value:input.value}); };
   td.append(input); tr.append(td);
  }); body.append(tr);
 });
}
search.oninput = render;
document.querySelector('#propose').onclick = () => parent.postMessage({type:'spike-data-edits', edits:[...edits.values()]}, '*');
render();
</script>'''.replace("PAYLOAD", payload)


def write_edits(path: Path, changes: list[dict]) -> None:
    """Prepare beside the source and replace only after serialization succeeds. Never evaluate user formulas."""
    book = None
    if path.suffix.lower() == ".csv":
        rows = table(path)
        for change in changes:
            rows[change["row"] + 1][change["column"]] = change["after"]
        stream = io.StringIO(newline="")
        csv.writer(stream).writerows(rows)
        blob = stream.getvalue().encode("utf-8")
    else:
        book = workbook(path)
        try:
            for change in changes:
                cell = book.worksheets[0].cell(change["row"] + 2, change["column"] + 1)
                if cell.data_type == "f":
                    raise ValueError("data_app_formula")
                value = change["after"]
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    number = float(value)
                    if (not math.isfinite(number) or (number == 0 and Decimal(value) != 0)
                            or len(Decimal(value).normalize().as_tuple().digits) > 15):
                        raise ValueError("data_app_value")
                    cell.value = number
                else:
                    cell.value = value
                    cell.data_type = "s"
            stream_bytes = io.BytesIO()
            book.save(stream_bytes)
            blob = stream_bytes.getvalue()
        finally:
            book.close()
    fd, name = tempfile.mkstemp(dir=path.parent, suffix=path.suffix)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(blob)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)
