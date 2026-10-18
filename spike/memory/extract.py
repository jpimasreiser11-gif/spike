"""Plain text from the files people put in a project, so memory can search them. Office and PDF files need the
optional `docs` libraries; anything else unreadable is skipped, never guessed."""

from __future__ import annotations

import html
import re
from pathlib import Path

from spike.memory.store import MAX_TEXT

TEXT = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".toml", ".xml",
        ".log", ".tex", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".c", ".h", ".cpp", ".sql",
        ".sh", ".ini", ".cfg"}
HTML = {".html", ".htm"}


def extract_text(path: Path) -> str | None:
    """The file's text (cut to the memory limit), or None when its format is not readable."""
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT:
            text = path.read_text(encoding="utf-8", errors="replace")
        elif suffix in HTML:
            raw = path.read_text(encoding="utf-8", errors="replace")
            raw = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", raw)
            text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
            text = re.sub(r"[ \t]+", " ", text)
        elif suffix == ".pdf":
            from pypdf import PdfReader

            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        elif suffix == ".docx":
            import docx

            doc = docx.Document(path)
            rows = ["\t".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
            text = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()] + rows)
        elif suffix == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            text = "\n\n".join(f"# {ws.title}\n" + "\n".join(
                "\t".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True))
                for ws in wb.worksheets)
        elif suffix == ".pptx":
            import pptx

            deck = pptx.Presentation(path)
            text = "\n\n".join("\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
                               for slide in deck.slides)
        else:
            return None
    except ImportError:
        return None
    except Exception:  # a damaged or encrypted file: nothing to index
        return None
    text = text.strip()
    return text[:MAX_TEXT] if text else None
